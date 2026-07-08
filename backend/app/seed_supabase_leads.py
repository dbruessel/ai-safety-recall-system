import os
import sys
import pandas as pd
import openpyxl
import requests
from dotenv import load_dotenv

# Load environment variables from local .env
load_dotenv()

# Setup Supabase Keys and fallback values
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY") or os.getenv("SUPABASE_SERVICE_KEY")

FILE_NAME = "clark-county-leads-v4.xlsx"

def discover_header_row(file_path):
    """
    Scans the first 30 rows of the spreadsheet to find the row containing 'Company Name'.
    This avoids hardcoding row offsets and makes the script highly resilient to layout changes.
    """
    wb = openpyxl.load_workbook(file_path, read_only=True)
    sheet = wb.active
    for r_idx, row in enumerate(sheet.iter_rows(max_row=30, values_only=True)):
        # Check if "Company Name" is in any cell of this row
        row_values = [str(cell).strip() if cell is not None else "" for cell in row]
        if "Company Name" in row_values:
            return r_idx
    return None

def main():
    print(f"📖 Reading spreadsheet from: {FILE_NAME}...")
    
    if not os.path.exists(FILE_NAME):
        print(f"❌ Error: Cannot locate {FILE_NAME}. Please download the spreadsheet first.")
        sys.exit(1)
        
    # Programmatically find headers
    header_row = discover_header_row(FILE_NAME)
    if header_row is None:
        print("❌ Error: Could not locate column headers (looking for 'Company Name') in the first 30 rows.")
        sys.exit(1)
        
    print(f"🔍 Programmatically identified column headers on Excel Row {header_row + 1}")
    
    # Read Excel using discovered header index
    df = pd.read_excel(FILE_NAME, header=header_row)
    
    # Clean up column names (strip whitespace)
    df.columns = [str(col).strip() for col in df.columns]
    
    # Check for core columns
    required_cols = ["Company Name", "Direct Corporate Email", "Est. Fleet Size"]
    for col in required_cols:
        if col not in df.columns:
            print(f"❌ Missing expected column in sheet: {col}")
            print(f"Available columns: {list(df.columns)}")
            sys.exit(1)

    # Clean text columns to avoid whitespace issues and silence the Pandas object dtype warning
    for col in df.columns:
        if pd.api.types.is_object_dtype(df[col]) or pd.api.types.is_string_dtype(df[col]):
            df[col] = df[col].astype(str).str.strip()

    # Drop any row where either Company Name or Direct Corporate Email is missing
    df = df.dropna(subset=["Company Name", "Direct Corporate Email"])
    df = df[df["Direct Corporate Email"] != "nan"]
    df = df[df["Company Name"] != "nan"]

    # --- CRITICAL FIX: DEDUPLICATE IN PYTHON BEFORE POSTING ---
    # Postgres ON CONFLICT DO UPDATE cannot process duplicate keys in a single batch statement.
    initial_count = len(df)
    df = df.drop_duplicates(subset=["Direct Corporate Email"], keep="first")
    deduped_count = len(df)
    
    if initial_count > deduped_count:
        print(f"🧹 Deduplicated contact emails: Removed {initial_count - deduped_count} duplicate rows in memory.")

    print(f"✅ Successfully parsed {len(df)} unique prospect leads from the spreadsheet.")

    # Convert dataframe into API-ready JSON format
    leads_to_insert = []
    for _, row in df.iterrows():
        # Get float USDOT OOS rate, handle non-numbers
        try:
            raw_rate = row.get("USDOT OOS Rate") or row.get("USDOT OOS %")
            if pd.isna(raw_rate) or raw_rate == "nan" or raw_rate == "":
                oos_rate = None
            else:
                # Convert "24.5%" or 24.5 to float
                rate_str = str(raw_rate).replace("%", "").strip()
                oos_rate = float(rate_str)
        except Exception:
            oos_rate = None

        # Fleet size robust conversion
        try:
            fleet_size = int(float(row["Est. Fleet Size"])) if pd.notna(row["Est. Fleet Size"]) and str(row["Est. Fleet Size"]).strip() != "" else 0
        except Exception:
            fleet_size = 0

        # Build payload mapping
        lead_payload = {
            "company_name": str(row["Company Name"]),
            "industry": str(row.get("Industry / NAICS", "Other")),
            "est_fleet_size": fleet_size,
            "primary_vehicle_mix": str(row.get("Primary Vehicle Mix", "")) if pd.notna(row.get("Primary Vehicle Mix")) else None,
            "contact_name": str(row.get("Primary Decision-Maker", "")) if pd.notna(row.get("Primary Decision-Maker")) else None,
            "contact_email": str(row["Direct Corporate Email"]),
            "contact_phone": str(row.get("Direct Phone Line", "")) if pd.notna(row.get("Direct Phone Line")) else None,
            "usdot_number": str(row.get("USDOT Number", "")) if pd.notna(row.get("USDOT Number")) else None,
            "localized_threat_hook": str(row.get("Localized Threat Hook", "")) if pd.notna(row.get("Localized Threat Hook")) else None,
            "lead_status": str(row.get("Lead Status", "Warm Prospect"))
        }
        
        # Strip string fields of standard 'nan' placeholders
        for key, val in lead_payload.items():
            if val == "nan" or val == "":
                lead_payload[key] = None

        leads_to_insert.append(lead_payload)

    # Check for Credentials
    if not SUPABASE_URL or not SUPABASE_KEY:
        print("\n⚠️ WARNING: SUPABASE_URL or SUPABASE_SERVICE_KEY/SUPABASE_SERVICE_ROLE_KEY environment variables not found.")
        print("🤖 Running in DRY RUN mode. All parsed entries are structurally valid and ready to ingest.")
        if leads_to_insert:
            print("Here's a sample of the first 2 lead records that will be inserted:\n")
            for idx, item in enumerate(leads_to_insert[:2]):
                print(f"Record {idx + 1}:")
                for k, v in item.items():
                    print(f"  {k}: {v}")
                print()
        print("💡 Ready to load! Once you have configured your .env file, run: python seed_supabase_leads.py")
        sys.exit(0)

    print("🔓 Credentials found! Connecting to Supabase...")
    
    # Target endpoint
    url = f"{SUPABASE_URL.rstrip('/')}/rest/v1/leads"
    
    # Configuration Headers for Bulk Insert with Upsert capabilities (On Conflict Merge)
    headers = {
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {SUPABASE_KEY}",
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates"  # Ask PostgREST to merge on conflict
    }

    # Stream leads in chunked batches of 100 to stay well under API payload thresholds
    batch_size = 100
    total_leads = len(leads_to_insert)
    print(f"🚀 Streaming {total_leads} unique leads into Supabase with ON CONFLICT DO UPDATE (Deduplication enabled)...")

    for i in range(0, total_leads, batch_size):
        batch = leads_to_insert[i:i + batch_size]
        
        # Append on_conflict parameter to tell PostgREST which unique constraint to match against
        batch_url = f"{url}?on_conflict=contact_email"
        
        try:
            response = requests.post(batch_url, json=batch, headers=headers)
            if response.status_code not in [200, 201]:
                print(f"❌ API Error in batch {i//batch_size + 1}: {response.status_code} - {response.text}")
                print("💡 Tip: If you see schema/column issues, check your database table structure.")
                sys.exit(1)
            else:
                end_range = min(i + batch_size, total_leads)
                print(f"  Uploaded leads {i + 1} to {end_range}...")
        except Exception as e:
            print(f"❌ Connection error in batch {i//batch_size + 1}: {e}")
            sys.exit(1)

    print("🎉 Ingestion completed successfully! Your leads table is live, synced, and deduplicated.")

if __name__ == "__main__":
    main()
